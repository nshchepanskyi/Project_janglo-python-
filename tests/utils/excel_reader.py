"""Read case details from the two source workbooks without losing source IDs."""
from pathlib import Path

from openpyxl import load_workbook


FIELDS = {
    "ID тест-кейсу": "id",
    "Назва тест-кейсу": "title",
    "Передумови": "preconditions",
    "Опис": "description",
}


def read_workbook(path):
    workbook = load_workbook(Path(path), data_only=True, read_only=True)
    cases = []
    for sheet in workbook.worksheets:
        values = {}
        header_row = None
        steps = []
        for row in sheet.iter_rows(values_only=True):
            cells = list(row)
            if len(cells) >= 4 and str(cells[0]).strip() == "Крок №":
                header_row = True
                continue
            if header_row and cells[0] is not None and str(cells[0]).strip().isdigit():
                steps.append({
                    "number": cells[0], "action": cells[1], "test_data": cells[2],
                    "expected": cells[3],
                })
            if len(cells) >= 2 and cells[0] in FIELDS:
                values[FIELDS[cells[0]]] = cells[1]
        if values.get("id"):
            cases.append({**values, "steps": steps, "sheet": sheet.title})
    return cases
